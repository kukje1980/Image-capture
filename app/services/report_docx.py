"""Word(.docx) 보고서 — 환자 정보 · 검사 · 영상 · 소견 · 서명란을 편집 가능한 문서로 만든다.

PDF/인쇄(app/ui/pro/print_util.py)와 같은 구성이지만, 진료실에서 소견을 고쳐 쓰거나
다른 문서에 붙여 넣을 수 있게 MS Word 형식으로 저장한다. Qt 없이 동작한다.

    build_docx(path, patient, items, exam=..., findings=..., per_page=2)
    items: [(bgr ndarray, caption str), ...]
"""
from __future__ import annotations

import io
from datetime import datetime

import numpy as np

try:
    import cv2
except ImportError:  # pragma: no cover
    cv2 = None

_SEX = {"M": "남", "F": "여", "O": "기타"}
_FONT = "Malgun Gothic"


def _png_stream(frame: np.ndarray) -> io.BytesIO:
    if cv2 is not None:
        ok, buf = cv2.imencode(".png", frame)
        if not ok:
            raise RuntimeError("영상을 PNG 로 바꾸지 못했습니다.")
        return io.BytesIO(buf.tobytes())
    from PIL import Image
    out = io.BytesIO()
    Image.fromarray(np.ascontiguousarray(frame[..., ::-1])).save(out, format="PNG")
    out.seek(0)
    return out


def _set_font(run, size=None, bold=None, color=None):
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    run.font.name = _FONT
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), _FONT)            # 한글 글꼴
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color.lstrip("#").upper())


def _para(doc_or_cell, text="", size=10.5, bold=False, color=None, align=None, space_after=2):
    from docx.shared import Pt
    p = doc_or_cell.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        _set_font(p.add_run(text), size, bold, color)
    return p


def _patient_rows(patient, exam, doctor: str) -> list[tuple[str, str]]:
    rows = []
    if patient is not None:
        rows += [("환자명", patient.name or "-"), ("차트번호", patient.chart_no or "-"),
                 ("성별", _SEX.get(patient.sex or "", "-")), ("생년월일", patient.birth_date or "-"),
                 ("진단명", patient.diagnosis or "-"), ("담당의", doctor or patient.doctor or "-")]
    if exam is not None:
        rows += [("검사일", exam.exam_date or "-"), ("검사", exam.title or "-")]
        if exam.diagnosis:
            rows.append(("검사 소견", exam.diagnosis))
    return rows


def build_docx(path: str, patient, items: list, exam=None, findings: str = "", doctor: str = "",
               title: str = "슬릿램프 영상 보고서", clinic: str = "", per_page: int = 2,
               rule_color: str = "#0d9488") -> int:
    """보고서를 path 에 저장하고 실린 영상 장수를 돌려준다."""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm

    if not items:
        return 0
    per_page = per_page if per_page in (1, 2, 4) else 2
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)       # A4
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(1.8))
    sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.6)
    usable = sec.page_width - sec.left_margin - sec.right_margin

    # --- 머리말 ---
    head = f"{clinic}  {title}" if clinic else title
    _para(doc, head, 17, True, "#0f172a", space_after=0)
    _para(doc, f"출력 {datetime.now():%Y-%m-%d %H:%M}", 8.5, False, "#64748b",
          WD_ALIGN_PARAGRAPH.RIGHT, space_after=0)
    _para(doc, "━" * 40, 8, False, rule_color, space_after=6)

    # --- 환자 · 검사 정보 (2열 표) ---
    rows = _patient_rows(patient, exam, doctor)
    if rows:
        pairs = [rows[i:i + 2] for i in range(0, len(rows), 2)]
        t = doc.add_table(rows=len(pairs), cols=4)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r, pair in enumerate(pairs):
            for c, (k, v) in enumerate(pair):
                kc, vc = t.cell(r, c * 2), t.cell(r, c * 2 + 1)
                kc.text = vc.text = ""
                _set_font(kc.paragraphs[0].add_run(k), 9.5, True, "#334155")
                _set_font(vc.paragraphs[0].add_run(str(v)), 9.5)
        _para(doc, space_after=4)

    # --- 영상 (per_page 장씩: 1 = 한 장 크게, 2 = 위아래, 4 = 2×2) ---
    pages = [items[i:i + per_page] for i in range(0, len(items), per_page)]
    for pno, page in enumerate(pages):
        if pno > 0:
            doc.add_page_break()
        cols = 2 if per_page == 4 else 1
        width = usable if cols == 1 else int(usable / 2) - Cm(0.3)
        if per_page == 2:
            width = int(usable * 0.78)
        grid = doc.add_table(rows=(len(page) + cols - 1) // cols, cols=cols)
        grid.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, (frame, caption) in enumerate(page):
            cell = grid.cell(i // cols, i % cols)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(_png_stream(frame), width=width)
            _para(cell, caption or "", 8.5, False, "#475569", WD_ALIGN_PARAGRAPH.CENTER, space_after=6)

    # --- 소견 · 서명 ---
    _para(doc, space_after=4)
    _para(doc, "소견", 11.5, True, "#0f172a", space_after=2)
    box = doc.add_table(rows=1, cols=1)
    box.style = "Table Grid"
    lines = (findings or "").splitlines() or [""]
    cell = box.cell(0, 0)
    cell.paragraphs[0].text = ""
    _set_font(cell.paragraphs[0].add_run(lines[0]), 10)
    for line in lines[1:]:
        _para(cell, line, 10, space_after=0)
    for _ in range(max(0, 4 - len(lines))):
        _para(cell, "", 10, space_after=0)
    _para(doc, space_after=10)
    sign = doctor or (patient.doctor if patient is not None and patient.doctor else "")
    _para(doc, f"판독의 / 확인자 :  {sign or '                    '}   (서명)",
          10, False, "#0f172a", WD_ALIGN_PARAGRAPH.RIGHT)
    doc.save(path)
    return len(items)
